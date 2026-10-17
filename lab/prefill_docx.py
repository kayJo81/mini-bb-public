"""
프리필 문서(prefill.json) → Word(.docx). 회의 중 직접 고쳐 쓰는 용도라 꾸밈은 최소로.
"""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Malgun Gothic"
INK, SUB, ACC = RGBColor(0x11, 0x18, 0x27), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x25, 0x63, 0xEB)
COLOR = {"up": RGBColor(0xC0, 0x39, 0x2B), "dn": RGBColor(0x1F, 0x62, 0xB5)}


def _font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _para(doc, text="", size=9.5, bold=False, color=None, italic=False, space=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.space_before = Pt(0)
    if text:
        _font(p.add_run(text), size, bold, color, italic)
    return p


def _shade(cell, hex_):
    tcpr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_)
    tcpr.append(sh)


def _cell(c, v, size=8.5, head=False):
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if isinstance(v, dict):
        t, b, col = v.get("t", ""), v.get("b", False), COLOR.get(v.get("c"))
    else:
        t, b, col = str(v), False, None
    _font(p.add_run(str(t)), size, bold=b or head, color=SUB if head else col)


def _table(doc, cols, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(cols):
        _cell(t.rows[0].cells[i], h, head=True)
        _shade(t.rows[0].cells[i], "F3F4F6")
    for r in rows:
        cs = t.add_row().cells
        for i, v in enumerate(r[:len(cols)]):
            _cell(cs[i], v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    _para(doc, space=1)
    return t


def write_docx(d, path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))

    _para(doc, "미니 블벅 프리필 회의록", 18, True, INK, space=0)
    _para(doc, f"{d['date'].replace('-', '.')} ({d['weekday']}) | 일간회의 | 자동 프리필 — 판단 칸은 회의에서 채움", 10, False, SUB, space=6)
    _table(doc, ["항목", "내용"], d["meta"], [3.5, 23])

    if d.get("summary"):
        _para(doc, "한눈에", 12, True, ACC, space=2)
        for s in d["summary"]:
            p = _para(doc, "", space=1)
            _font(p.add_run("• " + s), 9.5)
        _para(doc, space=4)

    for s in d["sections"]:
        _para(doc, s["title"], 14, True, ACC, space=4)
        for b in s["blocks"]:
            if b.get("title"):
                _para(doc, b["title"], 11, True, INK, space=2)
            if b.get("sub"):
                _para(doc, b["sub"], 9, False, SUB, italic=True, space=2)
            k = b["type"]
            if k == "table":
                _table(doc, b["cols"], b["rows"])
            elif k == "kv":
                _table(doc, ["항목", "내용"], b["rows"], [4, 22.5])
            elif k == "news":
                for i, it in enumerate(b["items"], 1):
                    p = _para(doc, "", space=0)
                    _font(p.add_run(f"{i}. {it['h']}"), 9.5, True)
                    if it.get("m"):
                        _para(doc, "   " + it["m"], 8.5, False, SUB, space=2)
                _para(doc, space=2)
            elif k == "para":
                _para(doc, b.get("text", ""), 9.5, space=4)
            elif k == "check":
                _table(doc, ["항목", "옵션"], [[{"t": a, "b": True}, "   ".join("▢ " + o for o in opts)] for a, opts in b["items"]], [4.5, 22])
            if b.get("note"):
                _para(doc, b["note"], 7.5, False, SUB, space=6)

    _para(doc, "회의 결과 기록", 12, True, ACC, space=2)
    _table(doc, ["항목", "결정"], [["레짐", ""], ["포지션", ""], ["무효화 조건", ""], ["다음 검토", ""], ["확신도", "    / 10"]], [4.5, 22])
    _para(doc, "투자 권유 아님 · 숫자는 자동 수집, 신호 기준은 각 표 아래 · 미니 블벅 랩", 7.5, False, SUB)
    doc.save(path)
